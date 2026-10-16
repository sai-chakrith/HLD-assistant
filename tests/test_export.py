import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.compare import compare
from app.export import NOTE, export_diff, export_findings, export_inventory, export_report
from tests.test_compare import _doc, _ent


@pytest.fixture
def data(conn):
    _doc(conn, "p1", "a", "1.0", [{"date": "d", "version": "1.0", "changed_by": "x", "description": "init"}])
    _doc(conn, "p1", "b", "2.0")
    _ent(conn, "p1", "a", "1.0", "components", "Alpha", description="old")
    _ent(conn, "p1", "b", "2.0", "components", "Alpha", description="new")
    _ent(conn, "p1", "b", "2.0", "components", "Beta", description="b")
    _ent(conn, "p1", "b", "2.0", "ports", "P", component="Beta", direction="provided")
    return conn


def test_inventory_csv_columns(data, tmp_path):
    paths = export_inventory("p1", "csv", tmp_path, doc_id="b", conn=data)
    assert len(paths) == 4
    df = pd.read_csv(next(p for p in paths if "components" in p.name))
    for col in ("review_status", "doc_id", "version", "page", "section"):
        assert col in df.columns
    assert set(df["name"]) == {"Alpha", "Beta"} and set(df["review_status"]) == {"unreviewed"}


def test_inventory_excel_note_and_sheets(data, tmp_path):
    (p,) = export_inventory("p1", "xlsx", tmp_path, doc_id="b", conn=data)
    wb = load_workbook(p)
    assert wb.sheetnames == ["components", "ports", "interfaces", "signals"]
    for ws in wb:
        assert ws["A1"].value == NOTE
        assert [c.value for c in ws[2]][-5:] == ["review_status", "doc_id", "version", "page", "section"]
    assert wb["components"]["A3"].value == "Alpha"


def test_diff_exports(data, tmp_path):
    d = compare("p1", "a", "1.0", "b", "2.0", conn=data)
    (j,) = export_diff(d, "json", tmp_path)
    payload = json.loads(j.read_text(encoding="utf-8"))
    assert payload["note"] == NOTE
    rows = payload["tables"]["components"]
    assert {r["change"] for r in rows} == {"added", "changed"}
    assert all("review_status" in r and "doc_id" in r and "page" in r for r in rows)
    assert payload["tables"]["revisions"][0]["side"] == "A"
    (x,) = export_diff(d, "xlsx", tmp_path)
    assert "revisions" in load_workbook(x).sheetnames


def test_findings_and_report(data, tmp_path):
    from app.checks import ensure_schema
    ensure_schema(data)
    data.execute("INSERT INTO findings(project_id, doc_id, version, rule_id, severity, title, explanation, "
                 "entities_json, evidence_json, label, review_status) VALUES('p1','b','2.0','M1','low','t','e',"
                 "'[\"Beta\"]','[{\"doc_id\":\"b\",\"version\":\"2.0\",\"page\":2,\"section\":\"1\"}]',"
                 "'potential','unreviewed')")
    data.commit()
    (p,) = export_findings("p1", "csv", tmp_path, conn=data)
    df = pd.read_csv(p)
    assert df.loc[0, "page"] == 2 and df.loc[0, "review_status"] == "unreviewed"
    (r,) = export_report({"rows": [{"component": "Beta", "doc_id": "b", "version": "2.0", "page": 1,
                                    "section": "1"}]}, "xlsx", tmp_path)
    assert load_workbook(r)["report"]["A1"].value == NOTE
    (r2,) = export_report({"title": "x", "citations": [{"doc_id": "b", "version": "2.0", "page": 3, "section": "2"}]},
                          "json", tmp_path)
    assert json.loads(r2.read_text())["tables"]["report"][0]["page"] == 3
