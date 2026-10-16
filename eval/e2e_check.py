"""End-to-end check of the whole system with the final model routing, through the real API (uvicorn + Ollama),
LLM cache reads off. Runs on copies of the DB / Chroma / processed data (eval/_work/e2e) so the demo data is not
changed; traces go to the real TRACE_PATH (data/traces.jsonl).

Steps: a cold start, b upload + timings, c rejected uploads, d six questions as viewer, e role checks,
f compare + exports, g cross-project isolation, h offline ask + extraction, i peak VRAM / `ollama ps`, then the 25
gold questions once more. Appends to eval/run.log and the "Final configuration" section of eval/results.md.

    .venv\\Scripts\\python -m eval.e2e_check
"""
from __future__ import annotations

import csv
import io
import json
import os
import shutil
import socket
import subprocess
import sys
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
_HF = ("HF_HUB_OFFLINE", "TRANSFORMERS_OFFLINE")
BASE_ENV = {k: v for k, v in os.environ.items() if k not in _HF}  # online unless step h

from eval.run_eval import (EVAL_DIR, WORK, GpuMonitor, _ollama, is_refusal, key_fact_hits, lat_stats,  # noqa: E402
                           ms, pct, table, trace_entries, unload_all)

SCRATCH = "e2e-scratch"
PDF_V1 = ROOT / "data/synthetic/hld_v1.pdf"
STEPS: list[dict] = []


def log(line: str) -> None:
    print(line, flush=True)
    with open(EVAL_DIR / "run.log", "a", encoding="utf-8") as f:
        f.write(line + "\n")


def step(name: str, ok: bool, result: str, seconds: float | None = None) -> None:
    STEPS.append({"step": name, "ok": bool(ok), "result": result, "seconds": seconds})
    log(f"[{'PASS' if ok else 'FAIL'}] {name}: {result}" + (f" ({seconds:.1f} s)" if seconds is not None else ""))


def loaded() -> list[str]:
    return [f"{m['name']} {100 * (m.get('size_vram') or 0) / max(1, m.get('size') or 1):.0f}% GPU / "
            f"{100 - 100 * (m.get('size_vram') or 0) / max(1, m.get('size') or 1):.0f}% CPU"
            for m in _ollama("/api/ps").get("models", [])]


class Server:
    def __init__(self, w: Path, offline: bool):
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            self.port = s.getsockname()[1]
        self.base = f"http://127.0.0.1:{self.port}"
        env = {**BASE_ENV, "DB_PATH": str(w / "hld.db"), "CHROMA_DIR": str(w / "chroma"),
               "PROCESSED_DIR": str(w / "processed"), "LLM_CACHE_READ": "0", "WARMUP": "1",
               "PYTHONIOENCODING": "utf-8"}
        if offline:
            env.update({k: "1" for k in _HF})
        self.t0 = time.perf_counter()
        self.proc = subprocess.Popen([sys.executable, "-m", "uvicorn", "app.api:app", "--port", str(self.port)],
                                     cwd=ROOT, env=env, stdout=open(w / f"uvicorn_{self.port}.log", "w"),
                                     stderr=subprocess.STDOUT)

    def wait_ready(self) -> tuple[float, dict]:
        import httpx
        while True:
            try:
                h = httpx.get(self.base + "/health", timeout=5).json()
                if h["warmup"]["status"] != "warming":
                    return time.perf_counter() - self.t0, h
            except httpx.HTTPError:
                pass
            if time.perf_counter() - self.t0 > 600 or self.proc.poll() is not None:
                raise SystemExit("server did not become ready")
            time.sleep(0.25)

    def stop(self) -> None:
        self.proc.terminate()
        self.proc.wait(timeout=30)


def prepare() -> Path:
    from app.config import settings
    w = WORK / "e2e"
    shutil.rmtree(w, ignore_errors=True)
    w.mkdir(parents=True)
    shutil.copy(settings.db_path, w / "hld.db")
    shutil.copytree(settings.chroma_dir, w / "chroma")
    shutil.copytree(settings.processed_dir, w / "processed")
    from app import controls
    c = controls.get_conn(w / "hld.db")
    controls.create_project(SCRATCH, "E2E scratch project", conn=c)
    for u, r in (("admin", "admin"), ("reviewer", "reviewer"), ("viewer", "viewer")):
        controls.add_member(SCRATCH, u, r, conn=c)
    c.close()
    return w


def _ask(base: str, pid: str, q: str, user: str = "viewer") -> dict:
    import httpx
    t = time.perf_counter()
    r = httpx.post(f"{base}/projects/{pid}/ask", headers={"X-User": user}, json={"question": q}, timeout=600)
    a = r.json()
    a["_ms"] = (time.perf_counter() - t) * 1000
    a["_status"] = r.status_code
    return a


def _cites(a: dict) -> str:
    return ", ".join(f"{c['doc_id']} p.{c['page_start']}" for c in a.get("citations", [])[:4]) or "-"


def _check_file(content: bytes, name: str, fmt: str, need: str = "review_status") -> tuple[bool, str]:
    """Opens the export (zip of several files, or one file) and checks the review_status column/key."""
    files = {}
    if content[:2] == b"PK" and fmt != "xlsx" or name.endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            files = {n: z.read(n) for n in z.namelist()}
    else:
        files = {name: content}
    ok, seen = True, []
    for n, b in files.items():
        if n.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True)
            heads = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows(max_row=1) for c in row]
            has = need in heads or any(need in str(c.value or "") for ws in wb.worksheets
                                       for row in ws.iter_rows(max_row=3) for c in row)
        elif n.endswith(".csv"):
            rows = list(csv.reader(io.StringIO(b.decode("utf-8-sig"))))
            has = bool(rows) and (need in rows[0] or any(need in ",".join(r) for r in rows[:3]))
        elif n.endswith(".json"):
            has = need in json.dumps(json.loads(b))
        else:
            has = False
        ok &= has
        seen.append(f"{Path(n).name}{'' if has else ' (no ' + need + ')'}")
    return ok, f"{len(files)} file(s): {', '.join(seen)}"


def main() -> int:
    import httpx
    from app.config import settings
    w = prepare()
    off = len(trace_entries(settings.trace_path))
    gpu = GpuMonitor(every=2.0)
    gpu.start()
    log(f"\n## e2e check {time.strftime('%Y-%m-%d %H:%M')}: Q&A/summaries {settings.demo_model} (keep_alive "
        f"{settings.demo_keep_alive}), extraction/upload {settings.llm_extract_model}; real API on copies "
        f"({w}); LLM cache reads off")
    out: dict = {"generated": time.strftime("%Y-%m-%d %H:%M")}
    A, R, V = {"X-User": "admin"}, {"X-User": "reviewer"}, {"X-User": "viewer"}
    # a. cold start
    unload_all()
    srv = Server(w, offline=False)
    try:
        secs, h = srv.wait_ready()
        out["cold_start_s"] = secs
        step("a. cold start (models unloaded, uvicorn start -> /health ready after warming the Q&A model)",
             h["warmup"]["status"] == "ready" and h["demo_model"] == settings.demo_model,
             f"warmup {h['warmup']['status']} steps {h['warmup']['steps']}; loaded: {loaded()}", secs)
        base = srv.base
        # b. upload approved synthetic v1 into the scratch project
        t = time.perf_counter()
        with open(PDF_V1, "rb") as f:
            r = httpx.post(f"{base}/projects/{SCRATCH}/upload", headers=A, timeout=3600,
                           files={"file": ("e2e_hld_v1.pdf", f, "application/pdf")},
                           data={"doc_type": "hld", "version": "1.0", "source_status": "approved"})
        up_s = time.perf_counter() - t
        up = r.json() if r.status_code == 200 else {"error": r.text[:300]}
        out["upload"] = up
        doc = up.get("doc_id")
        inv = httpx.get(f"{base}/projects/{SCRATCH}/inventory", headers=V, params={"doc_id": doc}, timeout=60).json() \
            if doc else {}
        fnd = httpx.get(f"{base}/projects/{SCRATCH}/findings", headers=V, params={"doc_id": doc}, timeout=60).json() \
            if doc else []
        n_ent = sum(len(v) for v in inv.values()) if isinstance(inv, dict) else 0
        tm = up.get("timings", {})
        out["loaded_after_extraction_before_rewarm"] = "see trace (re-warm runs inside the request)"
        out["loaded_after_upload"] = loaded()
        step("b. upload approved synthetic v1 (LLM extraction + checks + Q&A re-warm)",
             r.status_code == 200 and n_ent > 0 and len(fnd) > 0 and (up.get("rewarm") or {}).get("model")
             == settings.demo_model,
             f"HTTP {r.status_code}; ingest {ms(tm.get('ingest_ms'))}, extraction ({up.get('extract_model')}) "
             f"{ms(tm.get('extract_ms'))}, checks {ms(tm.get('checks_ms'))}, re-warm {ms(tm.get('rewarm_ms'))} "
             f"(load {ms((up.get('rewarm') or {}).get('load_ms'))}); entities {n_ent} ({up.get('extracted')}), "
             f"findings {len(fnd)}; loaded after: {out['loaded_after_upload']}", up_s)
        # c. rejected uploads
        t = time.perf_counter()
        res = []
        with open(PDF_V1, "rb") as f:
            raw = f.read()
        for label, fn, content, status in (("draft", "e2e_draft.pdf", raw + b"\n%draft\n", "draft"),
                                           ("non-PDF", "notes.txt", b"just text, not a pdf", "approved"),
                                           ("duplicate", "e2e_hld_v1_copy.pdf", raw, "approved")):
            rr = httpx.post(f"{base}/projects/{SCRATCH}/upload", headers=A, timeout=600,
                            files={"file": (fn, content, "application/octet-stream")},
                            data={"doc_type": "hld", "version": "1.1", "source_status": status})
            res.append((label, rr.status_code, rr.json().get("detail", "")[:70] if rr.status_code != 200 else "ACCEPTED"))
        step("c. draft / non-PDF / duplicate uploads rejected with a reason (400, audit-logged)",
             all(s == 400 for _, s, _ in res), "; ".join(f"{a}: {s} '{d}'" for a, s, d in res),
             time.perf_counter() - t)
        # d. six questions as viewer
        qs = [("ports of a named component", SCRATCH, "Which ports does the WinCtrlManager have?", "answer"),
              ("a section", SCRATCH, "What happens in the manual window operation flow?", "answer"),
              ("an identifier", SCRATCH, "Who sends the MotorCommand signal and who receives it?", "answer"),
              ("a specific release", "synthetic-ecu",
               "In release 1.0 of the Window Lift ECU HLD, which interfaces does the AntiPinchMonitor require?",
               "answer"),
              ("off-topic", SCRATCH, "What is a good recipe for pancakes with blueberries?", "off_topic"),
              ("on-topic but missing", SCRATCH, "What ASIL level is assigned to the WinCtrlManager?", "not_found")]
        rows, qa_ok = [], True
        for label, pid, q, want in qs:
            a = _ask(base, pid, q)
            ref = is_refusal(a.get("answer", ""), a.get("mode", ""))
            good = (not ref and bool(a.get("citations"))) if want == "answer" else \
                (a.get("mode") == "refused" and a.get("refusal_reason") == want)
            qa_ok &= good
            rows.append({"label": label, "question": q, "ok": good, "answer": a.get("answer", "")[:300],
                         "mode": a.get("mode"), "refusal_reason": a.get("refusal_reason"),
                         "confidence": a.get("confidence"), "model": a.get("model") or "-", "ms": a["_ms"],
                         "citations": _cites(a), "limitation": (a.get("limitation") or "")[:200]})
            log(f"   d {label}: {a.get('mode')} {a.get('refusal_reason') or ''} conf={a.get('confidence')} "
                f"model={a.get('model') or '-'} {a['_ms'] / 1000:.1f}s cites={_cites(a)} | {a.get('answer', '')[:160]}")
        out["questions"] = rows
        out["loaded_after_qa"] = loaded()
        step("d. six questions as viewer (3 answerable, 1 release, 1 off-topic refused, 1 not found)", qa_ok,
             f"{sum(r['ok'] for r in rows)}/6 as expected; latency p50 "
             f"{ms(lat_stats([r['ms'] for r in rows])['p50_ms'])}; models {sorted({r['model'] for r in rows})}; "
             f"loaded after Q&A: {out['loaded_after_qa']}", sum(r["ms"] for r in rows) / 1000)
        # e. roles
        t = time.perf_counter()
        fid = fnd[0]["id"] if fnd else 0
        codes = {
            "viewer approve": httpx.post(f"{base}/projects/{SCRATCH}/review", headers=V, timeout=60,
                                         json={"item_type": "finding", "item_id": fid, "decision": "approved"}).status_code,
            "viewer upload": httpx.post(f"{base}/projects/{SCRATCH}/upload", headers=V, timeout=60,
                                        files={"file": ("x.pdf", raw, "application/pdf")},
                                        data={"doc_type": "hld", "version": "9", "source_status": "approved"}).status_code,
            "viewer audit log": httpx.get(f"{base}/audit", headers=V, timeout=60).status_code,
            "reviewer approve": httpx.post(f"{base}/projects/{SCRATCH}/review", headers=R, timeout=60,
                                           json={"item_type": "finding", "item_id": fid, "decision": "approved",
                                                 "comment": "e2e check"}).status_code}
        step("e. viewer approve/upload/audit forbidden, reviewer approves a finding",
             codes["viewer approve"] == codes["viewer upload"] == codes["viewer audit log"] == 403
             and codes["reviewer approve"] == 200, ", ".join(f"{k} {v}" for k, v in codes.items()),
             time.perf_counter() - t)
        # f. compare two releases + exports
        t = time.perf_counter()
        cmp = httpx.post(f"{base}/projects/synthetic-ecu/compare", headers=V, timeout=300,
                         json={"doc_a": "syn_hld_v1", "doc_b": "syn_hld_v2"}).json()
        n_changes = sum(len(cmp.get(k, {}).get(x, [])) for k in ("components", "ports", "interfaces", "signals")
                        for x in ("added", "removed", "changed"))
        ex_res, ex_ok = [], n_changes > 0
        for kind, url, params in (("inventory", f"/projects/{SCRATCH}/export", {"kind": "inventory"}),
                                  ("findings", f"/projects/{SCRATCH}/export", {"kind": "findings"}),
                                  ("diff", "/projects/synthetic-ecu/export/diff",
                                   {"doc_a": "syn_hld_v1", "doc_b": "syn_hld_v2"})):
            pid_ = SCRATCH if kind != "diff" else "synthetic-ecu"
            for fmt in ("xlsx", "csv", "json"):
                rr = httpx.get(base + url, headers=V, params={**params, "fmt": fmt}, timeout=300)
                name = rr.headers.get("content-disposition", "").split("filename=")[-1].strip('"') or f"x.{fmt}"
                ok, info = _check_file(rr.content, name, fmt) if rr.status_code == 200 else (False, f"HTTP {rr.status_code}")
                if kind == "findings" and fmt == "json" and ok:
                    ok = "approved" in rr.text  # the finding approved in step e carries its status
                    info += "" if ok else " (approved status missing)"
                ex_ok &= ok
                ex_res.append(f"{kind}.{fmt} {'ok' if ok else 'FAIL'}")
                log(f"   f {pid_} {kind}.{fmt}: {info}")
        step("f. compare v1->v2 and export inventory/findings/diff as xlsx/csv/json (open + review_status)", ex_ok,
             f"diff entity changes {n_changes} ({str(cmp.get('summary', ''))[:120]}); {', '.join(ex_res)}",
             time.perf_counter() - t)
        # g. cross-project
        t = time.perf_counter()
        g = _ask(base, "chassis", "Which component drives the window lift motor and which signal does it receive?")
        chassis_docs = {d["doc_id"] for d in httpx.get(f"{base}/projects/chassis/documents", headers=V,
                                                       timeout=60).json()}
        leak = [c["doc_id"] for c in g.get("citations", []) if c["doc_id"] not in chassis_docs] + \
            [x for x in ("MotorDriver", "MotorCommand", "WinCtrlManager") if x in g.get("answer", "")]
        step("g. ask in chassis about content that exists only in synthetic-ecu/e2e-scratch (no leak)", not leak,
             f"{g.get('mode')} {g.get('refusal_reason') or ''}; citations {_cites(g)}; leaked: {leak or 'none'}",
             time.perf_counter() - t)
    finally:
        srv.stop()
    # h. offline restart + ask + extraction
    srv = Server(w, offline=True)
    try:
        secs, h = srv.wait_ready()
        base = srv.base
        a = _ask(base, SCRATCH, "Who sends the MotorCommand signal and who receives it?")
        t = time.perf_counter()
        r = httpx.post(f"{base}/projects/{SCRATCH}/extract", headers=A, timeout=3600,
                       json={"doc_id": doc, "use_llm": True})
        ex = r.json() if r.status_code == 200 else {}
        out["loaded_after_extraction"] = loaded()
        step("h. offline (HF_HUB_OFFLINE=1, TRANSFORMERS_OFFLINE=1): restart, one ask, one LLM re-extraction",
             h["warmup"]["status"] == "ready" and a.get("mode") == "generated" and r.status_code == 200,
             f"ready in {secs:.1f}s; ask {a.get('mode')} by {a.get('model')} in {a['_ms'] / 1000:.1f}s; re-extract "
             f"HTTP {r.status_code} {ex.get('extracted')} extraction {ms((ex.get('timings') or {}).get('extract_ms'))}"
             f", re-warm {ms((ex.get('timings') or {}).get('rewarm_ms'))}", time.perf_counter() - t + a["_ms"] / 1000)
        # 6. gold questions with the final routing
        gq = json.loads((EVAL_DIR / "gold_questions.json").read_text(encoding="utf-8"))["questions"]
        grows = []
        for q in gq:
            ga = _ask(base, q["project_id"], q["question"], user="admin")
            ref = is_refusal(ga.get("answer", ""), ga.get("mode", ""))
            okq = ref if q["category"] == "unanswerable" else (not ref and key_fact_hits(ga["answer"], q["key_facts"])
                                                                [0] == len(q["key_facts"]))
            grows.append({"id": q["id"], "category": q["category"], "correct": okq, "ms": ga["_ms"],
                          "model": ga.get("model"), "confidence": ga.get("confidence")})
        ans = [x for x in grows if x["category"] != "unanswerable"]
        un = [x for x in grows if x["category"] == "unanswerable"]
        out["gold"] = {"correctness": sum(x["correct"] for x in ans) / len(ans),
                       "refusal_accuracy": sum(x["correct"] for x in un) / len(un),
                       "latency": lat_stats([x["ms"] for x in grows]),
                       "models": sorted({x["model"] for x in grows if x["model"]}),
                       "wrong": [x["id"] for x in grows if not x["correct"]], "rows": grows}
        log(f"gold re-run: correct {pct(out['gold']['correctness'])}, refusal {pct(out['gold']['refusal_accuracy'])}, "
            f"wrong {out['gold']['wrong']}, models {out['gold']['models']}")
    finally:
        srv.stop()
    out["vram"] = gpu.stop()
    tr = trace_entries(settings.trace_path)[off:]
    out["swaps"] = [{"from": e.get("from_model"), "to": e["model"], "ms": e["latency_ms"], "ts": e["ts"]}
                    for e in tr if e.get("purpose") == "model_swap"]
    out["warm_calls"] = [{"purpose": e["purpose"], "model": e["model"], "ms": e["latency_ms"],
                          "load_ms": e.get("load_ms")} for e in tr if e.get("purpose") in ("warmup", "rewarm")]
    # chunks of the uploaded document in the scratch collection
    os.environ["CHROMA_DIR"] = str(w / "chroma")
    import chromadb
    from chromadb.config import Settings as CS
    from app.retrieve import collection_name
    cl = chromadb.PersistentClient(path=str(w / "chroma"), settings=CS(anonymized_telemetry=False))
    try:
        n_chunks = len(cl.get_collection(collection_name(SCRATCH, "bge")).get(where={"doc_id": doc})["ids"])
    except Exception as e:  # noqa: BLE001
        n_chunks, out["chunk_error"] = 0, str(e)
    out["chunks"] = n_chunks
    step("b'. chunks of the uploaded document indexed in the scratch project's collection", n_chunks > 0,
         f"{n_chunks} chunks")
    step("i. peak VRAM during the whole run (nvidia-smi every 2 s)", True,
         f"peak {out['vram'].get('peak_mib')} of {out['vram'].get('total_mib')} MiB; swaps "
         + "; ".join(f"{s['from'] or 'none'} -> {s['to']} {s['ms'] / 1000:.1f}s" for s in out["swaps"]))
    out["steps"] = STEPS
    R = json.loads((EVAL_DIR / "results.json").read_text(encoding="utf-8"))
    R["final_configuration"] = {**out, "demo_model": settings.demo_model, "extract_model": settings.llm_extract_model}
    (EVAL_DIR / "results.json").write_text(json.dumps(R, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    md = (EVAL_DIR / "results.md").read_text(encoding="utf-8").split("\n## Final configuration", 1)[0].rstrip()
    (EVAL_DIR / "results.md").write_text(md + "\n\n" + "\n".join(render(out, settings)) + "\n", encoding="utf-8")
    print("\n" + table(["step", "result", "time", "ok"], [[s["step"], s["result"][:180],
                                                          f"{s['seconds']:.1f} s" if s["seconds"] is not None else "-",
                                                          "PASS" if s["ok"] else "FAIL"] for s in STEPS]))
    return 0 if all(s["ok"] for s in STEPS) else 1


def render(o: dict, settings) -> list[str]:
    g = o.get("gold", {})
    L = ["## Final configuration (author's decision after seeing the comparison)", "",
         f"- Q&A answers, component reports and document summaries: **{settings.demo_model}** (DEMO_MODEL): "
         "95% vs 75% gold Q&A correctness in block C.",
         f"- LLM prose extraction and everything run at upload time: **{settings.llm_extract_model}** "
         "(LLM_EXTRACT_MODEL): faster, fewer false defects, fits fully in the 4 GB GPU; extraction runs once per "
         "upload, not while someone waits for an answer.",
         "- This decision was taken after seeing the numbers. The pre-set rule above favoured qwen2.5:3b for "
         "everything because qwen3:4b-instruct does not fit >= 90% on the GPU.",
         f"- Routing by purpose (qa/summary/report -> DEMO_MODEL; extract/upload -> LLM_EXTRACT_MODEL); overrides "
         f"only from ALLOWED_MODELS; think=false on every call; the Q&A model has keep_alive "
         f"{settings.demo_keep_alive}, is the only model warmed at start-up and is re-warmed after upload "
         "extraction; model loads are traced as 'model_swap' in data/traces.jsonl.", "",
         f"### End-to-end run ({o['generated']}; real API, LLM cache reads off, copies of the demo data)", "",
         table(["step", "result", "time", "ok"],
               [[s["step"], s["result"].replace("|", "/")[:400],
                 f"{s['seconds']:.1f} s" if s["seconds"] is not None else "-", "PASS" if s["ok"] else "FAIL"]
                for s in o["steps"]]), "",
         "Questions (step d):", "",
         table(["question", "mode", "confidence", "model", "latency", "citations", "answer"],
               [[f"{r['label']}: {r['question']}", f"{r['mode']} {r['refusal_reason'] or ''}".strip(),
                 r["confidence"], r["model"], ms(r["ms"]), r["citations"],
                 (r["answer"] or "")[:160].replace("|", "/")] for r in o.get("questions", [])]), "",
         "Model loads / swaps (data/traces.jsonl, Ollama load_duration):", "",
         table(["from", "to", "load time"], [[s["from"] or "none", s["to"], ms(s["ms"])] for s in o["swaps"]])
         if o["swaps"] else "none", "",
         f"Placement: after Q&A {', '.join(o.get('loaded_after_qa', [])) or '-'}; after the upload (incl. re-warm) "
         f"{', '.join(o.get('loaded_after_upload', [])) or '-'}; after the offline re-extraction (incl. re-warm) "
         f"{', '.join(o.get('loaded_after_extraction', [])) or '-'}. Peak GPU memory {o['vram'].get('peak_mib')} of "
         f"{o['vram'].get('total_mib')} MiB (includes ~1.1 GB used by other processes).", "",
         f"### Gold questions with the final routing (25, {', '.join(g.get('models', []))})", "",
         table(["metric", "block C (qwen3, eval harness)", "final configuration (API)"],
               [["answer correctness", "95%", pct(g.get("correctness"))],
                ["refusal accuracy", "100%", pct(g.get("refusal_accuracy"))],
                ["end-to-end latency p50 / p95", "5.1 s / 9.3 s",
                 f"{ms((g.get('latency') or {}).get('p50_ms'))} / {ms((g.get('latency') or {}).get('p95_ms'))}"],
                ["wrong", "R05", ", ".join(g.get("wrong", [])) or "none"]]), ""]
    return L


if __name__ == "__main__":
    sys.exit(main())
